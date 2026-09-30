"""
Báo cáo cho Quản trị nghiệp vụ (D1-12, TC-RP-01..04).

- summary(): số câu hỏi theo ngày/tuần/tháng, phân theo kết quả xử lý, số lần
  chuyển tư vấn viên và các câu hỏi phổ biến nhất.
- build_unanswered_workbook(): file Excel câu hỏi AI chưa trả lời được và câu
  hỏi đã chuyển tư vấn viên -> dùng để bổ sung FAQ.

Khoảng thời gian luôn là [date_from 00:00, date_to 24:00) theo giờ của MySQL.
"""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PERIODS = {
    "day": ("DATE_FORMAT(m.created_at, '%%Y-%%m-%%d')", "Ngày"),
    "week": ("DATE_FORMAT(m.created_at, '%%x-W%%v')", "Tuần"),
    "month": ("DATE_FORMAT(m.created_at, '%%Y-%%m')", "Tháng"),
}
# answer_status của tin trả lời -> tên cột trong báo cáo.
OUTCOMES = ("answered", "faq", "clarify", "cannot_answer", "out_of_scope", "blocked")
OUTCOME_LABELS = {
    "answered": "AI trả lời",
    "faq": "Trả lời FAQ",
    "clarify": "Hỏi lại",
    "cannot_answer": "Chưa trả lời được",
    "out_of_scope": "Ngoài phạm vi",
    "blocked": "Bị chặn",
}
MAX_RANGE_DAYS = 366
EXPORT_ROW_LIMIT = 5000


class ReportRangeError(ValueError):
    """Tham số khoảng thời gian không hợp lệ (thông báo hiển thị được cho người dùng)."""


def _parse_date(raw: str) -> date:
    try:
        return date.fromisoformat(raw)
    except ValueError:
        raise ReportRangeError("Ngày không hợp lệ, dùng định dạng YYYY-MM-DD.") from None


def parse_range(raw_from: str | None, raw_to: str | None, today: date | None = None):
    """Mặc định 30 ngày gần nhất. Trả về (date_from, date_to)."""
    today = today or date.today()
    date_to = _parse_date(raw_to) if raw_to else today
    date_from = _parse_date(raw_from) if raw_from else date_to - timedelta(days=29)
    if date_from > date_to:
        raise ReportRangeError("'from' phải nhỏ hơn hoặc bằng 'to'.")
    if (date_to - date_from).days >= MAX_RANGE_DAYS:
        raise ReportRangeError(f"Khoảng thời gian tối đa {MAX_RANGE_DAYS} ngày.")
    return date_from, date_to


def bucket_keys(period: str, date_from: date, date_to: date) -> list[str]:
    """Mọi kỳ trong khoảng (khớp định dạng của PERIODS) để kỳ không có dữ liệu vẫn
    hiện giá trị 0 trên biểu đồ thay vì bị bỏ qua."""
    keys, day = [], date_from
    while day <= date_to:
        if period == "day":
            key = day.isoformat()
        elif period == "week":
            iso_year, iso_week, _ = day.isocalendar()
            key = f"{iso_year}-W{iso_week:02d}"
        else:
            key = f"{day:%Y-%m}"
        if not keys or keys[-1] != key:
            keys.append(key)
        day += timedelta(days=1)
    return keys


def _empty_bucket(key: str) -> dict:
    return {"period": key, "questions": 0, **{outcome: 0 for outcome in OUTCOMES}}


def _bounds(date_from: date, date_to: date):
    return date_from.isoformat(), (date_to + timedelta(days=1)).isoformat()


def summary(cursor, period: str, date_from: date, date_to: date) -> dict:
    bucket_sql, _ = PERIODS[period]
    start, end = _bounds(date_from, date_to)

    cursor.execute(
        f"SELECT {bucket_sql} AS bucket, COUNT(*) AS total FROM messages m "
        "WHERE m.sender_type = 'user' AND m.created_at >= %s AND m.created_at < %s "
        "GROUP BY bucket ORDER BY bucket",
        (start, end),
    )
    series = {key: _empty_bucket(key) for key in bucket_keys(period, date_from, date_to)}
    for row in cursor.fetchall():
        series.setdefault(row["bucket"], _empty_bucket(row["bucket"]))["questions"] = row["total"]

    cursor.execute(
        f"SELECT {bucket_sql} AS bucket, m.answer_status AS outcome, COUNT(*) AS total "
        "FROM messages m WHERE m.sender_type IN ('model', 'system') "
        "AND m.answer_status IN (" + ", ".join(["%s"] * len(OUTCOMES)) + ") "
        "AND m.created_at >= %s AND m.created_at < %s GROUP BY bucket, outcome",
        (*OUTCOMES, start, end),
    )
    for row in cursor.fetchall():
        series.setdefault(row["bucket"], _empty_bucket(row["bucket"]))[row["outcome"]] = row["total"]

    cursor.execute(
        "SELECT COUNT(*) AS total FROM agent_notifications "
        "WHERE created_at >= %s AND created_at < %s",
        (start, end),
    )
    handovers = cursor.fetchone()["total"]

    cursor.execute(
        "SELECT MIN(m.content) AS question, COUNT(*) AS total FROM messages m "
        "WHERE m.sender_type = 'user' AND m.content <> '' "
        "AND m.created_at >= %s AND m.created_at < %s "
        "GROUP BY LOWER(TRIM(m.content)) ORDER BY total DESC LIMIT 10",
        (start, end),
    )
    top_questions = cursor.fetchall()

    rows = [series[key] for key in sorted(series)]
    totals = {"questions": sum(r["questions"] for r in rows), "handovers": handovers,
              **{outcome: sum(r[outcome] for r in rows) for outcome in OUTCOMES}}
    return {
        "period": period,
        "from": date_from.isoformat(),
        "to": date_to.isoformat(),
        "series": rows,
        "totals": totals,
        "top_questions": top_questions,
        # Danh sách (không phải dict) để giữ đúng thứ tự cột khi trả JSON.
        "outcomes": [{"key": key, "label": OUTCOME_LABELS[key]} for key in OUTCOMES],
    }


def fetch_unanswered(cursor, date_from: date, date_to: date) -> list[dict]:
    """Câu hỏi mà AI trả lời "chưa có thông tin" (cannot_answer)."""
    start, end = _bounds(date_from, date_to)
    cursor.execute(
        "SELECT m.created_at, m.conversation_id, c.channel, m.content AS reply, "
        "  (SELECT u.content FROM messages u WHERE u.conversation_id = m.conversation_id "
        "   AND u.sender_type = 'user' AND u.id < m.id ORDER BY u.id DESC LIMIT 1) AS question "
        "FROM messages m JOIN conversations c ON c.id = m.conversation_id "
        "WHERE m.answer_status = 'cannot_answer' AND m.created_at >= %s AND m.created_at < %s "
        "ORDER BY m.id LIMIT %s",
        (start, end, EXPORT_ROW_LIMIT),
    )
    return cursor.fetchall()


def fetch_handovers(cursor, date_from: date, date_to: date) -> list[dict]:
    """Các lần chuyển tư vấn viên kèm câu hỏi cuối cùng của thí sinh trước đó."""
    start, end = _bounds(date_from, date_to)
    cursor.execute(
        "SELECT n.created_at, n.conversation_id, c.channel, n.reason, c.status, "
        "  (SELECT u.content FROM messages u WHERE u.conversation_id = n.conversation_id "
        "   AND u.sender_type = 'user' AND u.created_at <= n.created_at "
        "   ORDER BY u.id DESC LIMIT 1) AS question "
        "FROM agent_notifications n JOIN conversations c ON c.id = n.conversation_id "
        "WHERE n.created_at >= %s AND n.created_at < %s ORDER BY n.id LIMIT %s",
        (start, end, EXPORT_ROW_LIMIT),
    )
    return cursor.fetchall()


def _write_sheet(sheet, headers, rows):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B2154")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        widest = max([len(str(header))] + [len(str(r[index - 1] or "")) for r in rows[:200]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(12, widest + 2), 70)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_unanswered_workbook(unanswered: list[dict], handovers: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chưa trả lời được"
    _write_sheet(
        sheet,
        ["Thời gian", "Phiên chat", "Kênh", "Câu hỏi của thí sinh", "Phản hồi của AI"],
        [[r["created_at"], r["conversation_id"], r["channel"], r["question"] or "", r["reply"]]
         for r in unanswered],
    )
    _write_sheet(
        workbook.create_sheet("Chuyển tư vấn viên"),
        ["Thời gian", "Phiên chat", "Kênh", "Lý do", "Trạng thái phiên", "Câu hỏi gần nhất"],
        [[r["created_at"], r["conversation_id"], r["channel"], r["reason"], r["status"],
          r["question"] or ""] for r in handovers],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

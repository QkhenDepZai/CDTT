import io
import unicodedata

import pytest

from knowledge.document_loader import (
    DocumentLoadError,
    decode_text_bytes,
    detect_source_type,
    load_document_bytes,
    normalize_text,
)


def _minimal_pdf(text: str) -> bytes:
    """Dựng 1 file PDF 1 trang hợp lệ có lớp text (không cần thư viện ngoài)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
              f"startxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_detect_source_type():
    assert detect_source_type("The_le.PDF") == "pdf"
    assert detect_source_type("faq.csv") == "csv"
    with pytest.raises(DocumentLoadError):
        detect_source_type("virus.exe")


def test_normalize_converts_nfd_to_nfc_and_cleans_spaces():
    decomposed = unicodedata.normalize("NFD", "Thể   lệ\r\n\n\n\ncuộc thi")
    assert normalize_text(decomposed) == "Thể lệ\n\ncuộc thi"


def test_decode_vietnamese_windows_1258():
    # cp1258 lưu "ệ" dạng tổ hợp: "ê" + dấu nặng (U+0323).
    raw = "L\u00ea\u0323 ph\u00ed".encode("cp1258")
    assert normalize_text(decode_text_bytes(raw)) == "Lệ phí"


def test_decode_utf16_with_bom():
    assert decode_text_bytes("Bảng A".encode("utf-16")) == "Bảng A"


def test_load_txt():
    sections = load_document_bytes("Thể lệ cuộc thi\n\nBảng A".encode("utf-8"), "a.txt")
    assert sections[0].text == "Thể lệ cuộc thi\n\nBảng A"


def test_load_csv_semicolon_includes_header_per_row():
    raw = "Bảng;Độ tuổi;Lệ phí\nA;12-18;Miễn phí\nB;19-24;\n".encode("utf-8-sig")
    sections = load_document_bytes(raw, "bang_thi.csv")
    assert [s.text for s in sections] == [
        "Bảng: A | Độ tuổi: 12-18 | Lệ phí: Miễn phí",
        "Bảng: B | Độ tuổi: 19-24",
    ]
    assert sections[0].metadata == {"row_number": 2}


def test_load_docx_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_heading("Quy chế thi", level=1)
    document.add_paragraph("Thí sinh có mặt trước 15 phút.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Bảng"
    table.cell(0, 1).text = "Thời gian"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "90 phút"
    buffer = io.BytesIO()
    document.save(buffer)

    text = load_document_bytes(buffer.getvalue(), "quy_che.docx")[0].text
    assert "Quy chế thi" in text
    assert "Thí sinh có mặt trước 15 phút." in text
    assert "A | 90 phút" in text


def test_load_pdf_with_text_layer():
    sections = load_document_bytes(_minimal_pdf("Python Master 2026 rules"), "rules.pdf")
    assert sections[0].page_number == 1
    assert "Python Master 2026 rules" in sections[0].text


def test_corrupted_pdf_raises_clear_error():
    with pytest.raises(DocumentLoadError):
        load_document_bytes(b"%PDF-1.4 not really a pdf", "broken.pdf")


def test_empty_file_raises():
    with pytest.raises(DocumentLoadError):
        load_document_bytes(b"", "empty.txt")
